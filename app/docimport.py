"""文档批量导入与自动切分（PRD M3-2）。

流程：解析文本 → 按标题/段落切分 → 合并过短片段、拆分过长片段 →
      生成「标准问」（取标题行，无标题则取首句）→ 先入草稿供人工预览 →
      人工确认后再发布（enabled=1）参与检索。

支持格式：粘贴文本 / .txt / .md / .docx（.docx 依赖 python-docx，未安装时给出明确提示）。
扫描版 PDF 需先转文本，本 MVP 不做 OCR。
"""
from __future__ import annotations

import re
from typing import Any

from . import config, db, knowledge

_HEADING = re.compile(r"^\s*(#{1,6}\s*|[一二三四五六七八九十]+[、.]\s*|\d+[、.]\s*)")

# 生成关键词时需要过滤的通用词
_STOP_TERMS = {
    "我们", "你们", "可以", "如果", "以及", "或者", "进行", "以下", "如下", "注意",
    "说明", "情况", "问题", "相关", "需要", "支持", "是否", "什么", "怎么", "如何",
    "政策", "规则", "条款", "内容", "信息",
}


def _clean_heading(line: str) -> str:
    return _HEADING.sub("", line).strip("# 　\t")


def _is_title_line(line: str, next_line: str = "") -> bool:
    """判断一行是否是「小节标题」。

    除了 markdown `#` / 数字序号，也要能识别纯文本里的小标题
    （如「七天无理由退货」这种独立成行、不超过 22 字、后面跟着正文的行），
    否则粘贴的纯文本文档会被整篇并成一个片段。
    """
    s = line.strip()
    if not s:
        return False
    if _HEADING.match(line):
        return True
    if len(s) > 22:
        return False
    if s.endswith(("：", ":")):
        return True
    if s.endswith(("。", "，", "、", "；", "！", "？")):
        return False
    nxt = (next_line or "").strip()
    return bool(nxt) and len(nxt) >= len(s)


def _looks_titled(block: str) -> bool:
    """片段的首行是否为小节标题（用于避免跨小节合并）。"""
    lines = block.split("\n")
    if len(lines) < 2:
        return bool(_HEADING.match(lines[0].strip()))
    return _is_title_line(lines[0], lines[1])


def split_text(text: str) -> list[str]:
    """按标题行与空行切分，再按长度规整为 200–600 字的片段。"""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")

    blocks: list[str] = []
    buf: list[str] = []
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            if buf:
                blocks.append("\n".join(buf).strip())
                buf = []
            continue
        next_line = ""
        for j in range(idx + 1, len(lines)):
            if lines[j].strip():
                next_line = lines[j]
                break
        if buf and _is_title_line(line, next_line):
            blocks.append("\n".join(buf).strip())
            buf = [stripped]
            continue
        buf.append(line.rstrip())
    if buf:
        blocks.append("\n".join(buf).strip())

    # 过短片段向后合并（但不跨越标题边界，否则整篇会被并成一个片段）
    merged: list[str] = []
    for block in blocks:
        if merged and len(merged[-1]) < config.CHUNK_MIN_CHARS and not _looks_titled(block):
            merged[-1] = f"{merged[-1]}\n{block}"
        else:
            merged.append(block)

    # 只有标题没有正文的片段（通常是文档大标题）丢弃，避免产生噪音条目
    def _title_only(block: str) -> bool:
        s = block.strip()
        return (len(block.split("\n")) == 1 and 0 < len(s) <= 22
                and not s.endswith(("。", "！", "？", "，", "、", "；")))

    merged = [b for b in merged if not _title_only(b)]

    # 过长片段按句子边界切分
    chunks: list[str] = []
    for block in merged:
        if len(block) <= config.CHUNK_MAX_CHARS:
            chunks.append(block)
            continue
        sentences = re.split(r"(?<=[。！？；!?;])", block)
        current = ""
        for sentence in sentences:
            if len(current) + len(sentence) > config.CHUNK_MAX_CHARS and current:
                chunks.append(current.strip())
                current = sentence
            else:
                current += sentence
        if current.strip():
            chunks.append(current.strip())

    return [c for c in chunks if len(c) >= 10]


def _make_question(chunk: str) -> str:
    first = chunk.split("\n", 1)[0].strip()
    first = _clean_heading(first)
    if not first:
        first = chunk.strip()[:40]
    if len(first) > 42:
        first = first[:42].rstrip("，。、;；:：") + "…"
    if not first.endswith(("？", "?", "…")):
        first += "？"
    return first


def _key_clauses(body: str, limit: int = 3, max_len: int = 30) -> list[str]:
    """从正文里挑出前几个短句，作为切片的「相似问法」匹配面。

    纯词面检索引擎需要足够的文字重叠才有召回，仅靠自动生成的标题往往不够，
    因此把正文的关键短句补充为相似问法。
    """
    clauses: list[str] = []
    for sentence in re.split(r"[。！？；\n]", body):
        sentence = sentence.strip(" ，。！？；、")
        if len(sentence) < 6:
            continue
        clauses.append(sentence[:max_len])
        if len(clauses) >= limit:
            break
    return clauses


def _make_keywords(question: str, clauses: list[str] | None = None) -> list[str]:
    terms: list[str] = []
    for piece in re.split(r"[，。、；：,.;:！？!?\s（）()【】\[\]—\-]+", question):
        piece = piece.strip("？?…")
        if 2 <= len(piece) <= 8 and piece not in _STOP_TERMS:
            terms.append(piece)
    for clause in clauses or []:
        for piece in re.split(r"[，。、；：,.;:！？!?\s（）()【】\[\]—\-]+", clause):
            piece = piece.strip("？?…")
            # 只保留 3–6 字的中文短语，过滤「99」「元需」这类噪声
            if (3 <= len(piece) <= 6 and piece not in _STOP_TERMS and piece not in terms
                    and not piece.isdigit()):
                terms.append(piece)
    return terms[:5]


def build_chunks(
    content: str, title: str, category: str, source_name: str = ""
) -> list[dict[str, Any]]:
    """把文档正文切分为知识条目草稿（不落库）。"""
    chunks = split_text(content)
    items: list[dict[str, Any]] = []
    for idx, chunk in enumerate(chunks, 1):
        question = _make_question(chunk)
        answer = chunk
        # 若首行是标题，答案里去掉重复的标题行，避免机器人复述标题
        lines = chunk.split("\n")
        if len(lines) > 1 and _clean_heading(lines[0]) and len(lines[0]) <= 42:
            answer = "\n".join(lines[1:]).strip() or chunk
        items.append(
            {
                "question": question,
                "answer": answer,
                "category": category or title or "文档导入",
                "similar": _key_clauses(answer),
                "keywords": _make_keywords(question, _key_clauses(answer, limit=2)),
                "enabled": False,
                "chunk_index": idx,
                "source_ref": f"{title or source_name} #{idx}",
                "source_name": source_name,
            }
        )
    return items


def import_document(
    title: str,
    content: str,
    category: str = "",
    source_name: str = "",
    source_type: str = "text",
    publish: bool = False,
) -> dict[str, Any]:
    """创建文档 + 切片。默认入草稿（不参与检索），由人工预览后发布。"""
    items = build_chunks(content, title, category, source_name)
    if not items:
        return {"document_id": None, "chunk_count": 0, "items": []}

    doc_id = db.create_document(title or source_name or "未命名文档", source_name, source_type)
    for item in items:
        item["doc_id"] = doc_id
        item["enabled"] = bool(publish)
        db.insert_knowledge(item)
    db.update_document(doc_id, chunk_count=len(items),
                       status="published" if publish else "draft")
    return {
        "document_id": doc_id,
        "chunk_count": len(items),
        "status": "published" if publish else "draft",
        "items": db.list_document_chunks(doc_id),
    }


def extract_upload(filename: str, data: bytes) -> tuple[str, str]:
    """从上传文件提取纯文本。返回 (文本, 类型)；不支持时抛 ValueError。"""
    lower = (filename or "").lower()
    if lower.endswith((".txt", ".md", ".markdown", ".csv")):
        for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
            try:
                return data.decode(encoding), "text"
            except UnicodeDecodeError:
                continue
        raise ValueError("无法识别文件编码，请另存为 UTF-8 后重试")

    if lower.endswith(".docx"):
        try:
            import docx  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise ValueError(
                "解析 .docx 需要 python-docx：pip install python-docx；或直接粘贴文本"
            ) from exc
        import io

        document = docx.Document(io.BytesIO(data))
        lines: list[str] = []
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(c for c in style if c.isdigit()) or "2"
                lines.append("#" * min(int(level), 6) + " " + text)
            else:
                lines.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines), "docx"

    if lower.endswith(".doc"):
        raise ValueError(".doc 为二进制老格式，请另存为 .docx 或 .txt 后重试")

    if lower.endswith(".pdf"):
        raise ValueError("当前 MVP 不支持 PDF（含扫描件），请复制文本后粘贴导入")

    raise ValueError("仅支持 .txt / .md / .docx / 直接粘贴文本")


def search_in_document(doc_id: int, query: str, top_k: int = 5) -> list[dict[str, Any]]:
    """文档发布前的切片检索预检：确认导入内容能被正确召回。"""
    chunks = db.list_knowledge(doc_id=doc_id)
    for chunk in chunks:
        chunk["enabled"] = True
    return knowledge.retrieve(query, chunks, top_k=top_k)
